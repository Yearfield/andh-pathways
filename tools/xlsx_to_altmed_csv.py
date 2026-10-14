#!/usr/bin/env python3
"""One-shot import: Vir_Harm_BVWP_Rosetta_Search workbook -> altmed-src/*.csv

    python3 -I tools/xlsx_to_altmed_csv.py <workbook.xlsx> altmed-src

After this import the CSV files are the source of truth: review and edit those, then run
tools/build_altmed.py. Do NOT re-run this import over reviewed CSVs (ids are regenerated).

Decisions baked in (agreed with the data owner):
  * only positive facts are stored; a feature that is not listed is treated as "no" (the app says so)
  * the header-less column 102 of the Formula sheet is all 1s (a leftover) and is ignored
  * obvious spelling / plural / spacing duplicates of a feature are merged (TYPOS + plural rule);
    near-matches (e.g. "ulcer" vs "mouth ulcer") are NEVER merged, they are listed for review
  * a species in both sheets becomes ONE entry: Rosetta text first, Formula strains folded in
"""
import sys, os, re, csv, collections, difflib
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
os.makedirs(os.path.join(OUT, "review"), exist_ok=True)
wb = openpyxl.load_workbook(SRC, data_only=True)
F = list(wb.worksheets[0].iter_rows(values_only=True))   # "Formula": one row per strain, prose + 0/1 features
M = list(wb.worksheets[1].iter_rows(values_only=True))   # "Rosetta Search": curated species, grouped 0/1 matrix


def clean(s):
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip() if s not in (None, "") else ""


# ---------- features: safe automatic merges ----------
TYPOS = {
    "enchephalitis": "encephalitis", "apendicitis": "appendicitis", "ptechiae": "petechiae", "hypotention": "hypotension",
    "erisypelas": "erysipelas", "arrythmia": "arrhythmia", "glosittis": "glossitis", "epiglotittis": "epiglottitis",
    "endocardits": "endocarditis", "endopthalmitis": "endophthalmitis", "allopecia": "alopecia", "cholecistitis": "cholecystitis",
    "pyelonefritis": "pyelonephritis", "rheumathoid arthritits": "rheumatoid arthritis", "sinusitus": "sinusitis",
    "nightsweats": "night sweats", "weightloss": "weight loss", "plattelets": "platelets", "painfull glands": "painful glands",
    "abcess": "abscess", "brain abcess": "brain abscess", "dental abcess": "dental abscess", "epididimytis": "epididymitis",
    "oopheritis": "oophoritis", "epididimo orchitis": "epididymo-orchitis", "broncho pneumonia": "bronchopneumonia",
    "deaf": "deafness", "non pathogenic": "non-pathogenic", "pid": "pelvic inflammatory disease", "chron s": "crohn s",
    "hemolytic": "haemolytic", "nosebleed": "epistaxis",
}


def fkey(label):
    k = re.sub(r"[^a-z0-9']+", " ", clean(label).lower().replace("'", " ")).strip()
    k = TYPOS.get(k, k)
    if len(k) > 3 and k.endswith("s") and not k.endswith(("ss", "is", "us", "ias")):
        k = k[:-1]
    return k


GROUP_TYPE = {"regions": "region", "Species": "host", "lifecycle": "stage", "Stage": "stage", "exposure": "exposure", "Route": "exposure",
              "Transmission": "exposure", "Trauma": "exposure", "Surface": "exposure", "tissue": "exposure", "System": "body_system",
              "Importance": "flag", "Appearance": "finding"}
feat = {}


def get_feat(label, group=None):
    k = fkey(label)
    if not k:
        return None
    f = feat.setdefault(k, {"label": k, "groups": collections.Counter(), "aliases": set()})
    if clean(label).lower() != k:
        f["aliases"].add(clean(label))
    if group:
        f["groups"][clean(group)] += 1
    return k


# ---------- pathogens ----------
KING = {"bacteria": "Bacteria", "nefbacteria": "Bacteria", "fungus": "Fungi", "fungi": "Fungi", "virus": "Virus", "worm": "Worm",
        "parasite": "Parasite", "amoeba": "Amoeba", "amoebozoa": "Amoeba", "protozoa": "Protozoa", "pathogen": "Other"}
RANKISH = re.compile(r"(ales|aceae|ineae|mycetes|cia|idae|viridae|virales)$", re.I)


def split_name(full):
    """'BACTERIA Actinobacteria Actinomycetales ... Actinomyces georgiae DSM 6843' -> ('Actinomyces georgiae DSM 6843', True).
    Falls back to the last few tokens (minus rank-looking words) when there is no 'Genus species' pair."""
    t = full.split()
    for i in range(1, len(t)):
        if t[i][:1].islower() and t[i].isalpha() and t[i - 1][:1].isupper():
            return " ".join(t[i - 1:]), True
    tail = t[-3:]
    while len(tail) > 1 and RANKISH.search(tail[0]):
        tail = tail[1:]
    return " ".join(tail) or full, False


paths = []
PROSE_F = {4: "vaccine", 5: "toxin", 6: "transmission", 7: "ab_resistance", 8: "host", 9: "commensal", 10: "disease", 11: "incubation",
           12: "symptom_a", 20: "risk_factors", 21: "fever", 22: "rash", 24: "cfs", 25: "toxin_details", 26: "aer_anaer", 27: "stain",
           28: "weightloss", 29: "diagnosis", 30: "treatment", 31: "prevention", 32: "bioweapon", 33: "size_um", 34: "notes",
           35: "notes_2", 36: "notes_3"}
SYMPTOM_COLS = [13, 14, 15, 16, 17, 18, 19, 23]
facts = set()
fh = F[3]
fcols = []
for j in range(37, len(fh)):
    fcols.append((j, get_feat(fh[j]) if fh[j] is not None else None))  # header-less columns (102 etc.) -> None -> ignored
for r in F[4:]:
    full = clean(r[3])
    disp, ok = split_name(full)
    p = {"src": "formula", "name": disp, "full_name": full, "kingdom": KING.get(full.split()[0].lower(), "Other") if full else "Other",
         "ncbi_taxid": clean(r[2]), "name_parsed": ok, "prose": {}, "strains": []}
    for j, n in PROSE_F.items():
        v = clean(r[j]) if j < len(r) else ""
        if v:
            p["prose"][n] = v
    sym = [clean(r[j]) for j in SYMPTOM_COLS if j < len(r) and clean(r[j])]
    if sym:
        p["prose"]["symptoms"] = " | ".join(sym)
    pi = len(paths)
    paths.append(p)
    for j, k in fcols:
        if k is not None and j < len(r) and r[j] == 1:
            facts.add((pi, k, "formula"))
mg, mf = M[0], M[1]
mcols = [(j, get_feat(mf[j], mg[j])) for j in range(4, len(mf)) if mf[j] is not None]
for r in M[3:]:
    if not r[0]:
        continue
    full = clean(r[0])
    disp, ok = split_name(full)
    p = {"src": "matrix", "name": disp, "full_name": full, "kingdom": KING.get(full.split()[0].lower(), "Other"), "ncbi_taxid": "",
         "name_parsed": ok, "prose": {}, "strains": []}
    if clean(r[2]):
        p["prose"]["host"] = clean(r[2])
    if clean(r[1]):
        p["prose"]["importance"] = clean(r[1])
    if clean(r[3]):
        p["prose"]["notes"] = clean(r[3])
    pi = len(paths)
    paths.append(p)
    for j, k in mcols:
        if j < len(r) and r[j] == 1:
            facts.add((pi, k, "matrix"))

# ---------- one entry per species: Rosetta text first, Formula strains folded in ----------
sk0 = lambda n: " ".join(n.lower().split()[:2])
fsp0 = collections.defaultdict(list)
for i, p in enumerate(paths):
    if p["src"] == "formula" and p["name_parsed"]:
        fsp0[sk0(p["name"])].append(i)
absorb = {}
for m, p in enumerate(paths):
    if p["src"] != "matrix" or not p["name_parsed"]:
        continue
    fl = fsp0.get(sk0(p["name"]), [])
    if not fl:
        continue
    p["strains"] = [(paths[i]["full_name"], paths[i]["ncbi_taxid"]) for i in fl]
    for i in fl:
        absorb[i] = m
        for n, v in paths[i]["prose"].items():
            if n not in p["prose"]:
                p["prose"][n] = v
            elif v not in p["prose"][n]:
                p["prose"][n + "_formula"] = v
        p["ncbi_taxid"] = p["ncbi_taxid"] or paths[i]["ncbi_taxid"]
    p["src"] = "matrix+formula"
facts = {(absorb.get(pi, pi), k, "matrix+formula" if pi in absorb else s) for pi, k, s in facts}
keep = [i for i in range(len(paths)) if i not in absorb]
remap = {o: n for n, o in enumerate(keep)}
facts = {(remap[pi], k, s) for pi, k, s in facts}
paths = [paths[i] for i in keep]

# ---------- ids ----------
for i, p in enumerate(paths):
    p["id"] = f"P{i + 1:04d}"
fkeys = sorted(feat)
for n, k in enumerate(fkeys):
    f = feat[k]
    f["id"] = f"F{n + 1:03d}"
    g = f["groups"].most_common()
    f["group"] = g[0][0] if g else "Ungrouped"
    f["type"] = GROUP_TYPE.get(f["group"], "clinical")
fmap = collections.defaultdict(set)
for pi, k, s in facts:
    fmap[(pi, k)].add(s)


def wcsv(path, header, rows):
    with open(os.path.join(OUT, path), "w", newline="", encoding="utf-8") as fh_:
        w = csv.writer(fh_)
        w.writerow(header)
        w.writerows(rows)


wcsv("features.csv", ["id", "label", "group", "type", "also_in_groups", "aliases"],
     [[feat[k]["id"], feat[k]["label"], feat[k]["group"], feat[k]["type"], ";".join(g for g, _ in feat[k]["groups"].most_common()[1:]),
       ";".join(sorted(feat[k]["aliases"]))] for k in fkeys])
wcsv("pathogens.csv", ["id", "name", "kingdom", "ncbi_taxid", "source_sheet", "full_name"],
     [[p["id"], p["name"], p["kingdom"], p["ncbi_taxid"], p["src"], p["full_name"]] for p in paths])
wcsv("facts.csv", ["pathogen_id", "feature_id", "relation", "source_sheet", "review_status"],
     [[paths[pi]["id"], feat[k]["id"], "present", "+".join(sorted(s)), "imported"] for (pi, k), s in sorted(fmap.items())])
wcsv("prose.csv", ["pathogen_id", "field", "text", "review_status"],
     [[p["id"], n, v, "imported"] for p in paths for n, v in p["prose"].items()])
wcsv("strains.csv", ["pathogen_id", "full_name", "ncbi_taxid"], [[p["id"], n, t] for p in paths for n, t in p["strains"]])

# ---------- review lists ----------
keys = list(feat)
sugg = []
for a in keys:
    for b in difflib.get_close_matches(a, keys, n=3, cutoff=0.86):
        if a < b:
            sugg.append((a, b))
    for b in keys:
        if a < b and len(a) > 3 and (a in b.split() or b in a.split()) and (a, b) not in sugg:
            sugg.append((a, b))
final = [["1_merged", feat[k]["id"], feat[k]["label"], feat[k]["group"], "absorbed: " + "; ".join(sorted(feat[k]["aliases"]))]
         for k in fkeys if feat[k]["aliases"]]
final += [["2_kept_separate", feat[a]["id"], a, feat[a]["group"], f"kept separate from: {b} ({feat[b]['id']})"] for a, b in sorted(sugg)]
wcsv("review/features_final_review.csv", ["section", "feature_id", "feature", "group", "note"], final)
wcsv("review/features_in_multiple_groups.csv", ["feature", "groups"],
     [[feat[k]["label"], ";".join(f"{g}:{c}" for g, c in feat[k]["groups"].items())] for k in fkeys if len(feat[k]["groups"]) > 1])
dups = collections.defaultdict(list)
for p in paths:
    dups[p["full_name"].lower()].append(p["id"])
wcsv("review/duplicate_pathogen_names.csv", ["full_name", "ids"], [[k, ";".join(v)] for k, v in dups.items() if len(v) > 1])
print(f"{len(paths)} pathogens, {len(feat)} features, {len(fmap)} positive facts, "
      f"{sum(1 for p in paths if p['src'] == 'matrix+formula')} merged species")
